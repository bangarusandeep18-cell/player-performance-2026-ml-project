import docx

doc = docx.Document('ML_Project_Report_Template (1).docx')
print(f"Tables count: {len(doc.tables)}")

for idx, tbl in enumerate(doc.tables):
    print(f"\n--- TABLE {idx} ---")
    print(f"Style: {tbl.style.name if tbl.style else 'None'}")
    print(f"Rows: {len(tbl.rows)}, Cols: {len(tbl.columns)}")
    for r_idx, row in enumerate(tbl.rows[:3]):
        cell_texts = [c.text.strip() for c in row.cells]
        print(f"  Row {r_idx}: {cell_texts}")
