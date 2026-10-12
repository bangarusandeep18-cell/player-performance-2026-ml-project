import docx

doc = docx.Document('ML_Project_Report_Template (1).docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

for i in range(min(35, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    font_names = set([r.font.name for r in p.runs if r.font.name])
    font_sizes = set([r.font.size.pt for r in p.runs if r.font.size])
    bolds = set([r.font.bold for r in p.runs])
    print(f"P{i:02d} | align={p.alignment} | style={p.style.name} | font={font_names} | size={font_sizes} | bold={bolds} | text='{p.text[:80]}'")
